from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import os
from django_rq import job

from stock.models import pick_strategy_result, pick_strategy
from stock.tasks import is_trade_day

from libs import wechat, constants


OUTPUT_DIR = '/data/share/log/stock/strategy'


@job('ths_worker', timeout=constants.JOB_TIMEOUT, result_ttl=constants.RESULT_TTL)
def generate_daily_summary(tx_date, ignore_trade_day=False):
    if not ignore_trade_day and not is_trade_day():
        return "非交易日"

    strategies = pick_strategy.objects.all()

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    title = doc.add_heading('选股汇总报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_para = doc.add_paragraph(f'日期：{tx_date}')
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    has_data = False

    for strategy in strategies:
        results = pick_strategy_result.objects.filter(
            strategy=strategy, pick_date=tx_date
        ).select_related('stock_code').order_by('stock_code__stock_code')

        if not results.exists():
            continue

        has_data = True

        doc.add_heading(strategy.strategy_name, level=1)

        info_keys = sorted(set(
            k for r in results for k in r.info_dict.keys()
        ))

        headers = ['代码', '名称'] + info_keys

        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Light Grid Accent 1'
        table.autofit = True

        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            for paragraph in hdr_cells[i].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)

        for r in results:
            row_cells = table.add_row().cells
            info = r.info_dict

            row_cells[0].text = r.stock_code.stock_code
            row_cells[1].text = r.stock_code.stock_name

            for j, key in enumerate(info_keys):
                val = info.get(key, '')
                row_cells[2 + j].text = str(val) if val is not None else ''

        doc.add_paragraph(f'共 {len(results)} 只股票').paragraph_format.space_before = Pt(4)

    if not has_data:
        doc.add_paragraph('当日无选股结果。')

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filepath = os.path.join(OUTPUT_DIR, f'{tx_date}.docx')
    doc.save(filepath)

    with open(filepath, 'rb') as f:
        media_id = wechat.upload_file(1, f.read(), 'file', f'{tx_date}.docx')

    if media_id:
        wechat.send_media_message(1, media_id, 'file')
        return f"✅ 汇总报告已保存: {filepath}"
    else:
        return f"❌ 文件上传失败: {filepath}"
